"""Выгрузка материалов в файлы: презентация в .pptx (её же открывает Canva), вопросы и идеи в .docx,
квиз в таблицу .xlsx для импорта в Kahoot.

Текст берётся из того же JSON, который учитель правит на странице, поэтому файл всегда
совпадает с тем, что он видел. Оформление простое и читаемое: шрифт, размеры, отступы.
"""
import io, re, zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.util import Cm as PCm, Pt as PPt

FONT = "Times New Roman"
SLIDE_FONT = "Arial"


def picture(images, item):
    """Картинка пункта из словаря {id: байты} или None."""
    return images.get(item.get("image")) if images and isinstance(item, dict) else None


def slides_pptx(title, slides, labels, images=None):
    """Презентация: титул, потом по слайду на пункт. Заметки учителя уходят в notes."""
    deck = Presentation()
    deck.slide_width, deck.slide_height = PCm(33.87), PCm(19.05)   # 16:9

    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = labels.get("subtitle", "")

    for item in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = item.get("title", "")
        body = slide.placeholders[1].text_frame
        body.clear()
        points = [p for p in item.get("points", []) if p.strip()]
        for i, point in enumerate(points):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = point
            para.font.size = PPt(20)
            para.font.name = SLIDE_FONT
        if item.get("question"):
            para = body.add_paragraph()
            para.text = f"{labels['question']}: {item['question']}"
            para.font.size = PPt(18)
            para.font.italic = True
            para.font.name = SLIDE_FONT
        if item.get("notes"):
            slide.notes_slide.notes_text_frame.text = item["notes"]
        image = picture(images, item)
        if image:  # текст слева, картинка справа, пропорции сохраняются
            holder = slide.placeholders[1]
            holder.width = int(deck.slide_width * 0.52) - holder.left
            try:
                pic = slide.shapes.add_picture(io.BytesIO(image), int(deck.slide_width * 0.55), PCm(4.5), height=PCm(12.5))
            except Exception:  # формат, который PowerPoint не берёт (например webp): слайд без картинки
                continue
            box = int(deck.slide_width * 0.42)
            if pic.width > box:
                pic.height, pic.width = int(pic.height * box / pic.width), box

    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def start_doc(title):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = head.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    return doc


def add_picture(doc, images, item):
    image = picture(images, item)
    if image:
        try:
            doc.add_picture(io.BytesIO(image), width=Cm(10))
        except Exception:  # формат, который Word не берёт: без картинки
            pass


def quiz_docx(title, questions, labels, images=None):
    """Вопросы для класса, ответы для учителя отдельным разделом в конце."""
    doc = start_doc(title)
    letters = "ABCDEFGH"
    for i, q in enumerate(questions, 1):
        para = doc.add_paragraph()
        para.add_run(f"{i}. ").bold = True
        para.add_run(q.get("question", ""))
        mark = para.add_run(f"  [{labels.get(q.get('level'), '')}]")
        mark.font.size = Pt(9)
        mark.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        add_picture(doc, images, q)
        for n, option in enumerate(q.get("options", [])):
            line = doc.add_paragraph(f"{letters[n % len(letters)]}) {option}")
            line.paragraph_format.left_indent = Cm(1)
        if not q.get("options"):
            doc.add_paragraph("_" * 60).paragraph_format.left_indent = Cm(1)

    doc.add_page_break()
    head = doc.add_paragraph()
    head.add_run(labels["answers"]).bold = True
    for i, q in enumerate(questions, 1):
        doc.add_paragraph(f"{i}. {q.get('answer', '')}")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def ideas_docx(title, ideas, labels, images=None):
    """Идеи заданий: по каждой шаги, что нужно, время и как проверить."""
    doc = start_doc(title)
    for i, idea in enumerate(ideas, 1):
        para = doc.add_paragraph()
        run = para.add_run(f"{i}. {idea.get('title', '')}")
        run.bold = True
        run.font.size = Pt(14)
        minutes = idea.get("minutes")
        if minutes:
            note = para.add_run(f"   {minutes} {labels['minutes']}")
            note.font.size = Pt(10)
            note.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        add_picture(doc, images, idea)
        doc.add_paragraph(idea.get("what", ""))
        for key, name in (("needs", labels["needs"]), ("assess", labels["assess"])):
            if idea.get(key):
                line = doc.add_paragraph()
                line.add_run(f"{name}: ").bold = True
                line.add_run(idea[key])
        doc.add_paragraph()
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def report_docx(title, about, parts):
    """Отчёт о проверке: заголовок работы, строка сведений и разделы с пунктами, пустые разделы пропускаются."""
    doc = start_doc(title)
    if about:
        doc.add_paragraph(about).alignment = WD_ALIGN_PARAGRAPH.CENTER
    for head, lines in parts:
        if not lines:
            continue
        run = doc.add_paragraph().add_run(head)
        run.bold = True
        run.font.size = Pt(14)
        for line in lines:
            doc.add_paragraph(line)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def xlsx(rows):
    """Одна таблица в .xlsx на zipfile: текст и числа, без стилей. Для импорта больше не нужно."""
    def cell(ref, value):
        if isinstance(value, int):
            return f'<c r="{ref}"><v>{value}</v></c>'
        text = escape(re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value)))  # управляющие символы ломают XML
        return f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{text}</t></is></c>'
    sheet = "".join(f'<row r="{i}">' + "".join(cell(f"{chr(65 + j)}{i}", v) for j, v in enumerate(row)) + "</row>"
                    for i, row in enumerate(rows, 1))
    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    parts = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '</Types>',
        "_rels/.rels": '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<?xml version="1.0" encoding="UTF-8"?><workbook {ns} xmlns:r="{rel}">'
            '<sheets><sheet name="Kahoot" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": f'<?xml version="1.0" encoding="UTF-8"?><worksheet {ns}><sheetData>{sheet}</sheetData></worksheet>',
    }
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, body in parts.items():
            z.writestr(name, body)
    return out.getvalue()


def nonum(text):
    """Без номера в начале строки: нумерацию ставит сам лист, иначе выходит «1. 1. Tere»."""
    return re.sub(r"^\s*\d{1,3}[.)]\s+", "", text)


KAHOOT_TIMES = (5, 10, 20, 30, 60, 90, 120, 240)   # другие значения Kahoot не принимает


def kahoot_xlsx(questions):
    """Таблица для «Import spreadsheet» в Kahoot: вопрос до 120 знаков, от двух до четырёх ответов до 75,
    время из разрешённого списка и номер правильного ответа. Вопрос без двух ответов пропускается."""
    rows = [["Question", "Answer 1", "Answer 2", "Answer 3", "Answer 4", "Time limit", "Correct answer(s)"]]
    for q in questions:
        given = [str(a).strip()[:75] for a in q.get("answers", [])][:4]
        answers = [a for a in given if a]
        if not str(q.get("question", "")).strip() or len(answers) < 2:
            continue
        # номер правильного ответа считаем после выброса пустых вариантов, иначе он съедет
        right = min(len(given), max(1, int(q.get("correct") or 1))) - 1
        correct = len([a for a in given[:right] if a]) + 1 if given[right] else 1
        seconds = min(KAHOOT_TIMES, key=lambda t: abs(t - int(q.get("seconds") or 20)))
        rows.append([q["question"].strip()[:120], *answers, *[""] * (4 - len(answers)), seconds, correct])
    return xlsx(rows)


def worksheet_docx(title, tasks, labels, images=None):
    """Рабочий лист для печати: варианты с новой страницы, настоящие таблицы и линии, ответы учителю в конце."""
    doc = start_doc(title)
    line = "_" * 64
    variant = None
    for t in tasks:
        if t.get("variant") and t["variant"] != variant:
            if variant is not None:
                doc.add_page_break()
            variant = t["variant"]
            head = doc.add_paragraph()
            head.alignment = WD_ALIGN_PARAGRAPH.CENTER
            head.add_run(variant).bold = True
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(12)
        run = para.add_run(t.get("title", ""))
        run.bold = True
        if t.get("score"):
            para.add_run(f"   ({t['score']} {labels['score']})").font.size = Pt(10)
        if t.get("instruction"):
            doc.add_paragraph().add_run(t["instruction"]).italic = True
        add_picture(doc, images, t)
        kind, items = t.get("kind"), [i for i in t.get("items", []) if i.strip()]
        if kind == "table":
            cells = lambda row: [c.strip() for c in row.split("|")]
            header, rows = cells(t.get("header", "")), [cells(r) for r in t.get("rows", []) if r.strip()]
            width = max([len(header)] + [len(r) for r in rows])
            table = doc.add_table(rows=1 + len(rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate([header] + rows):
                for c in range(width):
                    table.cell(r, c).text = row[c] if c < len(row) else ""
                    if r == 0:
                        for p in table.cell(r, c).paragraphs:
                            for x in p.runs:
                                x.bold = True
        elif kind == "questions":
            for n, q in enumerate(items, 1):
                doc.add_paragraph(f"{n}. {nonum(q)}")
                doc.add_paragraph(line)
        elif kind == "writing":
            for q in items:
                doc.add_paragraph(q)
            for _ in range(max(1, min(int(t.get("lines") or 8), 30))):
                doc.add_paragraph(line)
        else:  # gaps: строки с пропусками как есть
            for n, q in enumerate(items, 1):
                doc.add_paragraph(f"{n}. {nonum(q)}")
    keys = [t for t in tasks if t.get("answers")]
    if keys:
        doc.add_page_break()
        doc.add_paragraph().add_run(labels["answers"]).bold = True
        for t in keys:
            doc.add_paragraph(" ".join(filter(None, (t.get("variant"), t.get("title")))) + ": " + "; ".join(t["answers"]))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def build(kind, title, data, labels, images=None):
    """Файл материала: имя файла, тип и содержимое."""
    if kind == "worksheet":
        return "docx", ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        worksheet_docx(title, data.get("tasks", []), labels, images))
    if kind == "kahoot":
        return "xlsx", ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        kahoot_xlsx(data.get("questions", [])))
    if kind == "slides":
        return "pptx", ("application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        slides_pptx(title, data.get("slides", []), labels, images))
    if kind == "quiz":
        return "docx", ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        quiz_docx(title, data.get("questions", []), labels, images))
    return "docx", ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ideas_docx(title, data.get("ideas", []), labels, images))


if __name__ == "__main__":
    labels = {"question": "Вопрос классу", "subtitle": "Урок", "answers": "Ответы для учителя",
              "minutes": "мин", "needs": "Нужно", "assess": "Как проверить",
              "easy": "лёгкий", "medium": "средний", "hard": "трудный"}
    deck = slides_pptx("Тема урока", [{"title": "Первый слайд", "points": ["раз", "два"],
                                       "question": "почему?", "notes": "сказать вслух"}], labels)
    assert deck[:2] == b"PK" and len(deck) > 20000, len(deck)
    quiz = quiz_docx("Вопросы", [{"question": "Что это?", "kind": "choice", "options": ["а", "б"],
                                  "answer": "а", "level": "easy"}], labels)
    ideas = ideas_docx("Идеи", [{"title": "Аукцион", "what": "делают то-то", "needs": "доска",
                                 "minutes": 20, "assess": "сдают таблицу"}], labels)
    for blob in (quiz, ideas):
        assert blob[:2] == b"PK" and len(blob) > 10000, len(blob)
    kind, (mime, blob) = build("slides", "Тема", {"slides": []}, labels)
    assert kind == "pptx" and "presentation" in mime and blob[:2] == b"PK"
    table = kahoot_xlsx([{"question": "Что <это> & то?", "answers": ["а", "б" * 99, "", "в"], "correct": 4, "seconds": 27},
                         {"question": "Без ответов", "answers": ["один"], "correct": 1, "seconds": 20}])
    with zipfile.ZipFile(io.BytesIO(table)) as z:
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
    assert sheet.count("<row") == 2, sheet                  # вопрос с одним ответом пропущен
    assert "&lt;это&gt; &amp;" in sheet and "б" * 75 + "<" in sheet and "б" * 76 not in sheet
    assert '<c r="F2"><v>30</v></c>' in sheet and '<c r="G2"><v>3</v></c>' in sheet, sheet  # время к 30, ответ в пределах
    sheet = worksheet_docx("Kontrolltöö", [
        {"variant": "I variant", "title": "1. KÜSIMUSED", "kind": "questions", "items": ["Mis sinu nimi on?"], "score": 2},
        {"variant": "I variant", "title": "2. TEGUSÕNAD", "kind": "table", "header": "ISTUMA | MAGAMA",
         "rows": ["Ma istu____ | ", "Sa | "], "answers": ["istun, magan"]},
        {"variant": "II variant", "title": "1. KIRJUTA", "kind": "writing", "items": ["Mina"], "lines": 3}],
        {**labels, "score": "p"})
    with zipfile.ZipFile(io.BytesIO(sheet)) as z:
        body = z.read("word/document.xml").decode()
    assert body.count("<w:tbl>") == 1 and "Ma istu____" in body and "II variant" in body and "istun, magan" in body
    print("ok")
